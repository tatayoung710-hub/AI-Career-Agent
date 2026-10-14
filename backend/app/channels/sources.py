"""Authorized source readers with bounded network/file input and provenance."""

import csv
import hashlib
import io
import ipaddress
import json
import re
import socket
from urllib.parse import parse_qs, urlparse

import httpx
from defusedxml import ElementTree
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import settings
from app.models import Job, JobSourceRecord, JobVersion, Source, utcnow


def safe_url(url: str) -> str:
    parsed = urlparse(url)
    hostname = parsed.hostname or ""
    if parsed.scheme != "https" or parsed.username or parsed.password or parsed.port not in {None, 443}:
        raise ValueError("来源读取仅支持 HTTPS 标准端口且不含 URL 凭据")
    if hostname.lower() not in {h.lower() for h in settings.source_fetch_allowed_hosts}:
        raise ValueError("该来源主机尚未加入 SOURCE_FETCH_ALLOWED_HOSTS；可先上传表格或粘贴文章")
    for result in socket.getaddrinfo(hostname, 443, type=socket.SOCK_STREAM):
        address = ipaddress.ip_address(result[4][0])
        if not address.is_global:
            raise ValueError("禁止读取内网、回环或非公网地址")
    return url


def download(url: str, headers: dict | None = None) -> bytes:
    safe_url(url)
    with httpx.Client(timeout=20, follow_redirects=False) as client:
        with client.stream("GET", url, headers=headers) as response:
            response.raise_for_status()
            if response.is_redirect:
                raise ValueError("来源返回重定向，请使用授权的最终 URL")
            chunks, size = [], 0
            for chunk in response.iter_bytes():
                size += len(chunk)
                if size > settings.max_upload_bytes:
                    raise ValueError("来源内容超过上传大小限制")
                chunks.append(chunk)
    return b"".join(chunks)


def text_value(value) -> str:
    if isinstance(value, list):
        return " ".join(text_value(item) for item in value)
    if isinstance(value, dict):
        return str(value.get("link") or value.get("url") or value.get("text") or value.get("name") or "")
    return str(value or "")


def normalize_row(row: dict) -> dict:
    aliases = {
        "company": ["company", "公司", "公司名称", "企业", "企业名称"],
        "title": ["title", "岗位", "岗位名称", "职位", "职位名称", "招聘岗位"],
        "location": ["location", "地点", "城市", "工作地点", "工作城市"],
        "jd_text": ["jd_text", "jd", "JD", "要求", "岗位要求", "职位描述", "岗位描述", "招聘详情"],
        "source_url": ["source_url", "url", "链接", "投递链接", "招聘链接", "官网链接"],
        "salary_text": ["salary_text", "薪资", "薪酬"],
    }
    values = {}
    for field, keys in aliases.items():
        values[field] = next((text_value(row[key]) for key in keys if key in row and text_value(row[key]).strip()), "")
    values["jd_complete"] = bool(values["jd_text"].strip())
    if not values["jd_text"]:
        values["jd_text"] = "\n".join(f"{key}: {text_value(value)}" for key, value in row.items() if not key.startswith("_"))
    return values


def read_source(source: Source) -> list[dict]:
    if source.kind == "rss":
        root = ElementTree.fromstring(download(source.url))
        rows = []
        for item in root.findall(".//item")[:1000]:
            rows.append({"title": item.findtext("title") or "", "source_url": item.findtext("link") or "",
                         "jd_text": item.findtext("description") or "", "_external_id": item.findtext("guid") or ""})
        if not rows:
            ns = {"a": "http://www.w3.org/2005/Atom"}
            for item in root.findall("a:entry", ns)[:1000]:
                link = item.find("a:link", ns)
                rows.append({"title": item.findtext("a:title", namespaces=ns) or "",
                             "source_url": link.get("href", "") if link is not None else "",
                             "jd_text": item.findtext("a:summary", namespaces=ns) or item.findtext("a:content", namespaces=ns) or "",
                             "_external_id": item.findtext("a:id", namespaces=ns) or ""})
        return rows
    if source.kind in {"csv", "xlsx"}:
        data = download(source.url)
        if source.kind == "csv":
            return list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))[:1000]
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            header = [str(value or "") for value in next(rows, [])]
            return [dict(zip(header, values, strict=False)) for _, values in zip(range(1000), rows)]
        finally:
            workbook.close()
    if source.kind == "feishu" or ".feishu.cn/base/" in source.url:
        if not settings.feishu_access_token:
            raise ValueError("飞书同步需要配置 FEISHU_ACCESS_TOKEN 并授予目标多维表格读取权限")
        parsed = urlparse(source.url)
        app_token = parsed.path.rstrip("/").split("/")[-1]
        table_id = parse_qs(parsed.query).get("table", [""])[0]
        if not re.fullmatch(r"[A-Za-z0-9]+", app_token) or not re.fullmatch(r"tbl[A-Za-z0-9]+", table_id):
            raise ValueError("飞书来源 URL 必须包含 base app_token 和 table 参数")
        if app_token not in settings.source_feishu_allowed_bases:
            raise ValueError("该多维表格尚未获部署者授权：需加入 SOURCE_FEISHU_ALLOWED_BASES")
        rows, page_token = [], ""
        for _page in range(2):
            url = f"https://open.feishu.cn/open-apis/bitable/v1/apps/{app_token}/tables/{table_id}/records?page_size=500"
            if page_token:
                from urllib.parse import quote
                url += "&page_token=" + quote(page_token, safe="")
            result = json.loads(download(url, {"Authorization": f"Bearer {settings.feishu_access_token}"}))
            if result.get("code") != 0:
                raise ValueError(f"飞书读取失败，错误码 {result.get('code')}；检查读取权限与 token")
            for record in result.get("data", {}).get("items", []):
                rows.append({**record.get("fields", {}), "_external_id": record.get("record_id", "")})
            if not result.get("data", {}).get("has_more"):
                break
            page_token = result.get("data", {}).get("page_token", "")
        return rows
    raise ValueError("当前来源需要人工粘贴文章、JD 或上传 CSV/XLSX；尚未提供获准自动读取连接器")


def sync_source(db: Session, user_id: str, source_id: str) -> dict:
    source = db.scalar(select(Source).where(Source.id == source_id, Source.user_id == user_id, Source.archived.is_(False)))
    if not source:
        raise ValueError("来源不存在或已归档")
    rows = read_source(source)
    imported, skipped, errors = 0, 0, []
    for index, row in enumerate(rows):
        values = normalize_row(row)
        if not values["title"].strip():
            errors.append({"row": index + 1, "error": "缺少具体岗位名称，保留原来源后需人工补全"})
            continue
        external_key = str(row.get("_external_id") or values["source_url"] or hashlib.sha256(
            json.dumps(values, sort_keys=True, ensure_ascii=False).encode()).hexdigest())[:255]
        existing = db.scalar(select(JobSourceRecord.id).where(JobSourceRecord.user_id == user_id,
                                                              JobSourceRecord.source_key == source.id,
                                                              JobSourceRecord.external_key == external_key))
        if existing:
            skipped += 1
            continue
        job = Job(user_id=user_id, source_id=source.id, **values)
        db.add(job)
        db.flush()
        db.add(JobVersion(user_id=user_id, job_id=job.id, version_number=1,
                          jd_text=job.jd_text, requirements=[], snapshot=values, extraction_version="source-import-v1"))
        db.add(JobSourceRecord(user_id=user_id, job_id=job.id, source_id=source.id,
                               source_key=source.id, external_key=external_key,
                               url=values["source_url"], raw_text=job.jd_text))
        imported += 1
    source.last_synced_at, source.new_jobs_count = utcnow(), imported
    db.flush()
    return {"source_id": source.id, "imported": imported, "skipped": skipped, "errors": errors,
            "read_limit": 1000, "note": "同步只读取前 1000 条；字段缺失时请补全具体岗位和 JD。"}
