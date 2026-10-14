"""Owner-scoped career records, immutable versions, imports and real document exports."""

import csv
import hashlib
import io
import re
import tempfile
from itertools import islice
from datetime import date, datetime
from pathlib import Path
from typing import Annotated
from xml.sax.saxutils import escape
from zipfile import ZipFile

from docx import Document as WordDocument
from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile
from fastapi.encoders import jsonable_encoder
from fastapi.responses import FileResponse
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer
from sqlalchemy import func, inspect, or_, select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.ai.evidence import fact_snapshot, iter_claims, validate_resume_evidence
from app.auth import get_current_user
from app.config import settings
from app.db import get_db
from app.models import (
    Application,
    ApplicationEvent,
    Document,
    Fact,
    Job,
    JobSourceRecord,
    JobVersion,
    Profile,
    Resume,
    ResumeEvidence,
    ResumeVersion,
    Source,
    Task,
    User,
    new_id,
    utcnow,
)
from app.schemas import (
    ApplicationCreate,
    ApplicationUpdate,
    FactCreate,
    JobCreate,
    JobUpdate,
    ProfileUpdate,
    ResumeCreate,
    ResumeVersionCreate,
    SourceCreate,
    SourceUpdate,
)

router = APIRouter(tags=["Career records"])
DB = Annotated[Session, Depends(get_db)]
CurrentUser = Annotated[User, Depends(get_current_user)]


def serialize(record, *, exclude: tuple[str, ...] = ()) -> dict:
    return {
        column.key: getattr(record, column.key)
        for column in inspect(record).mapper.column_attrs
        if column.key not in exclude
    }


def owned(db: Session, model, entity_id: str, user_id: str):
    record = db.scalar(select(model).where(model.id == entity_id, model.user_id == user_id))
    if record is None:
        raise HTTPException(404, "记录不存在")
    return record


def commit(db: Session) -> None:
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "记录冲突，请刷新后重试") from None


def job_snapshot(job: Job) -> dict:
    return jsonable_encoder(serialize(job))


def create_job_record(
    db: Session, user_id: str, data: JobCreate, source_key: str = "manual"
) -> Job:
    if data.source_id:
        owned(db, Source, data.source_id, user_id)
    job = Job(user_id=user_id, **data.model_dump())
    db.add(job)
    db.flush()
    db.add(
        JobVersion(
            user_id=user_id,
            job_id=job.id,
            version_number=1,
            jd_text=job.jd_text,
            requirements=job.requirements,
            snapshot=job_snapshot(job),
        )
    )
    external_key = hashlib.sha256(
        f"{data.company}|{data.title}|{data.location}|{data.source_url}".encode()
    ).hexdigest()
    db.add(
        JobSourceRecord(
            user_id=user_id,
            job_id=job.id,
            source_id=data.source_id,
            source_key=data.source_id or source_key,
            external_key=external_key,
            url=data.source_url,
            raw_text=data.jd_text,
            published_at=data.published_at,
        )
    )
    return job


@router.get("/profile")
def get_profile(db: DB, user: CurrentUser) -> dict:
    profile = db.scalar(select(Profile).where(Profile.user_id == user.id))
    if profile is None:
        profile = Profile(user_id=user.id, email=user.email)
        db.add(profile)
        commit(db)
    return serialize(profile)


@router.patch("/profile")
def patch_profile(data: ProfileUpdate, db: DB, user: CurrentUser) -> dict:
    profile = db.scalar(select(Profile).where(Profile.user_id == user.id))
    if profile is None:
        profile = Profile(user_id=user.id, email=user.email)
        db.add(profile)
    for key, value in data.model_dump(exclude_unset=True, exclude_none=True).items():
        setattr(profile, key, value)
    profile.updated_at = utcnow()
    commit(db)
    return serialize(profile)


@router.get("/facts")
def list_facts(db: DB, user: CurrentUser, include_history: bool = False) -> list[dict]:
    query = select(Fact).where(Fact.user_id == user.id)
    if not include_history:
        superseded = select(Fact.supersedes_id).where(
            Fact.user_id == user.id, Fact.supersedes_id.is_not(None)
        )
        query = query.where(Fact.id.not_in(superseded))
    return [serialize(row) for row in db.scalars(query.order_by(Fact.created_at.desc()))]


@router.post("/facts", status_code=201)
def create_fact(data: FactCreate, db: DB, user: CurrentUser) -> dict:
    if data.document_id:
        owned(db, Document, data.document_id, user.id)
    values = data.model_dump()
    if data.supersedes_id:
        previous = owned(db, Fact, data.supersedes_id, user.id)
        already_revised = db.scalar(
            select(Fact.id).where(Fact.user_id == user.id, Fact.supersedes_id == previous.id)
        )
        if already_revised:
            raise HTTPException(409, "该事实已经有新修订，请从最新版本继续修改")
        values.update(lineage_id=previous.lineage_id, revision_number=previous.revision_number + 1)
    fact = Fact(user_id=user.id, **values)
    db.add(fact)
    commit(db)
    return serialize(fact)


@router.post("/facts/{fact_id}/confirm")
def confirm_fact(fact_id: str, db: DB, user: CurrentUser) -> dict:
    fact = owned(db, Fact, fact_id, user.id)
    if not fact.confirmed:
        fact.confirmed, fact.confirmed_at = True, utcnow()
        commit(db)
    return serialize(fact)


def decode_text(blob: bytes) -> str:
    if b"\x00" in blob:
        raise ValueError("文本文件包含二进制数据")
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return blob.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文本编码，请保存为 UTF-8 后上传")


def parse_document(blob: bytes, suffix: str) -> str:
    if suffix == ".pdf":
        if not blob.startswith(b"%PDF-"):
            raise ValueError("文件内容不是 PDF")
        pdf = PdfReader(io.BytesIO(blob))
        if pdf.is_encrypted:
            raise ValueError("请上传未加密 PDF")
        if len(pdf.pages) > 100:
            raise ValueError("简历 PDF 最多支持 100 页")
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        if not text.strip():
            raise ValueError("未找到可提取文字；扫描 PDF 请先 OCR，或上传 Word/TXT")
    elif suffix == ".docx":
        with ZipFile(io.BytesIO(blob)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > 100 * 1024 * 1024:
                raise ValueError("Word 文件解压体积超过限制")
            if "word/document.xml" not in archive.namelist():
                raise ValueError("文件内容不是 DOCX")
        doc = WordDocument(io.BytesIO(blob))
        parts = [paragraph.text for paragraph in doc.paragraphs]
        parts.extend(
            " | ".join(cell.text for cell in row.cells)
            for table in doc.tables
            for row in table.rows
        )
        text = "\n".join(parts)
    else:
        text = decode_text(blob)
    if not text.strip():
        raise ValueError("文件没有可读取内容")
    if len(text) > 300000:
        raise ValueError("提取文本超过 300,000 字符")
    return text


def public_document(document: Document) -> dict:
    return serialize(document, exclude=("storage_path",))


@router.get("/documents")
def list_documents(db: DB, user: CurrentUser) -> list[dict]:
    return [
        public_document(row)
        for row in db.scalars(
            select(Document).where(Document.user_id == user.id).order_by(Document.created_at.desc())
        )
    ]


@router.get("/documents/{document_id}")
def get_document(document_id: str, db: DB, user: CurrentUser) -> dict:
    return public_document(owned(db, Document, document_id, user.id))


@router.post("/documents", status_code=201)
def upload_document(db: DB, user: CurrentUser, file: UploadFile = File(...)) -> dict:
    filename = re.split(r"[/\\]", file.filename or "resume.txt")[-1][:255]
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pdf", ".docx", ".txt"}:
        raise HTTPException(415, "支持 PDF、DOCX、TXT；旧版 DOC 请先转换为 DOCX")
    blob = file.file.read(settings.max_upload_bytes + 1)
    if not blob or len(blob) > settings.max_upload_bytes:
        raise HTTPException(413, "文件为空或超过大小限制")
    document_id = new_id()
    folder = settings.storage_dir.resolve() / user.id / "documents"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{document_id}{suffix}"
    path.write_bytes(blob)
    document = Document(
        id=document_id,
        user_id=user.id,
        filename=filename,
        content_type=file.content_type or "",
        storage_path=str(path),
        sha256=hashlib.sha256(blob).hexdigest(),
        size_bytes=len(blob),
    )
    try:
        document.parsed_text = parse_document(blob, suffix)
        document.status = "parsed"
    except Exception as exc:
        document.status, document.error = "failed", str(exc)[:1000]
    db.add(document)
    commit(db)
    return public_document(document)


@router.get("/sources")
def list_sources(db: DB, user: CurrentUser) -> list[dict]:
    return [
        serialize(row)
        for row in db.scalars(
            select(Source)
            .where(Source.user_id == user.id, Source.archived.is_(False))
            .order_by(Source.created_at.desc())
        )
    ]


@router.post("/sources/catalog")
def add_source_catalog(db: DB, user: CurrentUser) -> dict:
    """Add the user-requested starter directory without adding fictional jobs or facts."""
    created = 0
    for name, kind, url, notes in SOURCE_ENTRIES:
        source = db.scalar(select(Source).where(Source.user_id == user.id, Source.name == name))
        if source is None:
            db.add(Source(user_id=user.id, name=name, kind=kind, url=url, notes=notes))
            created += 1
    commit(db)
    return {"created": created, "message": "来源入口已登记；登记不表示已经读取来源内容或具备提交权限。"}


@router.post("/sources", status_code=201)
def create_source(data: SourceCreate, db: DB, user: CurrentUser) -> dict:
    source = Source(user_id=user.id, **data.model_dump())
    db.add(source)
    commit(db)
    return serialize(source)


@router.patch("/sources/{source_id}")
def patch_source(source_id: str, data: SourceUpdate, db: DB, user: CurrentUser) -> dict:
    source = owned(db, Source, source_id, user.id)
    for key, value in data.model_dump(exclude_unset=True, exclude_none=True).items():
        setattr(source, key, value)
    commit(db)
    return serialize(source)


@router.delete("/sources/{source_id}", status_code=204)
def archive_source(source_id: str, db: DB, user: CurrentUser) -> Response:
    source = owned(db, Source, source_id, user.id)
    source.archived = True
    commit(db)
    return Response(status_code=204)


@router.post("/sources/{source_id}/sync", status_code=202)
def sync_source(source_id: str, db: DB, user: CurrentUser) -> dict:
    source = owned(db, Source, source_id, user.id)
    if source.archived:
        raise HTTPException(409, "来源已经归档")
    if source.sync_task_id:
        existing = db.scalar(
            select(Task).where(
                Task.id == source.sync_task_id,
                Task.user_id == user.id,
                Task.status.in_(["queued", "running", "retry"]),
            )
        )
        if existing:
            return serialize(existing)
    task = Task(
        user_id=user.id,
        kind="source_sync",
        input_json={"source_id": source.id},
        idempotency_key=f"source-sync:{source.id}:{new_id()}",
    )
    db.add(task)
    db.flush()
    source.sync_task_id = task.id
    commit(db)
    return serialize(task)


@router.get("/jobs")
def list_jobs(
    db: DB, user: CurrentUser, q: str = "", status: str | None = None, source_id: str | None = None
) -> list[dict]:
    query = select(Job).where(Job.user_id == user.id)
    if q:
        pattern = f"%{q}%"
        query = query.where(
            or_(
                Job.title.ilike(pattern),
                Job.company.ilike(pattern),
                Job.location.ilike(pattern),
                Job.jd_text.ilike(pattern),
            )
        )
    if status:
        query = query.where(Job.status == status)
    if source_id:
        query = query.where(Job.source_id == source_id)
    return [serialize(row) for row in db.scalars(query.order_by(Job.created_at.desc()))]


@router.post("/jobs", status_code=201)
def create_job(data: JobCreate, db: DB, user: CurrentUser) -> dict:
    job = create_job_record(db, user.id, data)
    commit(db)
    return serialize(job)


@router.get("/jobs/{job_id}")
def get_job(job_id: str, db: DB, user: CurrentUser) -> dict:
    return serialize(owned(db, Job, job_id, user.id))


@router.patch("/jobs/{job_id}")
def patch_job(job_id: str, data: JobUpdate, db: DB, user: CurrentUser) -> dict:
    job = owned(db, Job, job_id, user.id)
    if data.source_id:
        owned(db, Source, data.source_id, user.id)
    changes = data.model_dump(exclude_unset=True, exclude={"version_number"})
    if "jd_text" in changes and "jd_complete" not in changes:
        changes["jd_complete"] = True
    nullable = {"source_id", "published_at", "deadline"}
    if any(value is None and key not in nullable for key, value in changes.items()):
        raise HTTPException(422, "岗位文本字段不能设为 null")
    result = db.execute(
        update(Job)
        .where(Job.id == job.id, Job.user_id == user.id, Job.version_number == data.version_number)
        .values(**changes, version_number=data.version_number + 1, updated_at=utcnow())
    )
    if result.rowcount != 1:
        db.rollback()
        raise HTTPException(409, "岗位已更新，请刷新后重试")
    db.refresh(job)
    db.add(
        JobVersion(
            user_id=user.id,
            job_id=job.id,
            version_number=job.version_number,
            jd_text=job.jd_text,
            requirements=job.requirements,
            snapshot=job_snapshot(job),
        )
    )
    commit(db)
    return serialize(job)


IMPORT_ALIASES = {
    "company": ("company", "公司", "公司名称", "企业"),
    "title": ("title", "job", "岗位", "岗位名称", "职位", "职位名称"),
    "location": ("location", "地点", "城市", "工作地点"),
    "industry": ("industry", "行业"),
    "jd_text": ("jd_text", "jd", "JD", "要求", "岗位要求", "职位描述"),
    "salary_text": ("salary_text", "薪资", "工资"),
    "source_url": ("source_url", "url", "link", "链接", "投递链接", "招聘链接"),
    "published_at": ("published_at", "发布时间"),
    "deadline": ("deadline", "截止时间", "截止日期"),
}


def import_date(value) -> date | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value).strip().replace("/", "-")[:10])


@router.post("/jobs/import")
def import_jobs(
    db: DB,
    user: CurrentUser,
    file: UploadFile = File(...),
    preview: bool = True,
    source_id: str | None = None,
) -> dict:
    if source_id:
        owned(db, Source, source_id, user.id)
    blob = file.file.read(5 * 1024 * 1024 + 1)
    if len(blob) > 5 * 1024 * 1024:
        raise HTTPException(413, "岗位表格最多支持 5 MB")
    suffix = Path(file.filename or "").suffix.lower()
    try:
        if suffix == ".csv":
            raw_rows = list(islice(csv.DictReader(io.StringIO(decode_text(blob))), 1001))
        elif suffix == ".xlsx":
            with ZipFile(io.BytesIO(blob)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("表格解压体积超过限制")
            workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
            rows = workbook.active.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(rows)]
            raw_rows = [dict(zip(headers, row)) for row in islice(rows, 1001)]
            workbook.close()
        else:
            raise HTTPException(415, "岗位导入支持 CSV 或 XLSX")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f"无法读取表格：{exc}") from None
    if len(raw_rows) > 1000:
        raise HTTPException(422, "每次最多导入 1,000 行，请拆分表格")
    normalized, errors = [], []
    for index, row in enumerate(raw_rows, start=2):
        row = {str(key).strip(): value for key, value in row.items() if key is not None}
        if not any(value not in (None, "") for value in row.values()):
            continue
        values = {
            field: next((row[alias] for alias in aliases if row.get(alias) not in (None, "")), "")
            for field, aliases in IMPORT_ALIASES.items()
        }
        try:
            for field in ("published_at", "deadline"):
                values[field] = import_date(values[field])
            for field in set(values) - {"published_at", "deadline"}:
                values[field] = str(values[field]).strip()
            values["jd_complete"] = bool(values["jd_text"])
            values["jd_text"] = values["jd_text"] or "[待补充 JD：原表格仅提供招聘线索]"
            normalized.append(JobCreate(**values, source_id=source_id))
        except Exception as exc:
            errors.append({"row": index, "error": str(exc)[:500]})
    imported_count, skipped_count = 0, 0
    if not preview:
        for item in normalized:
            existing = db.scalar(
                select(Job.id).where(
                    Job.user_id == user.id,
                    Job.company == item.company,
                    Job.title == item.title,
                    Job.location == item.location,
                    Job.source_url == item.source_url,
                )
            )
            if existing:
                skipped_count += 1
                continue
            create_job_record(db, user.id, item, "import")
            imported_count += 1
        commit(db)
    return {
        "preview": preview,
        "rows": [item.model_dump() for item in normalized],
        "errors": errors,
        "imported_count": imported_count,
        "skipped_count": skipped_count,
    }


@router.get("/resumes")
def list_resumes(db: DB, user: CurrentUser) -> list[dict]:
    return [
        serialize(row)
        for row in db.scalars(
            select(Resume).where(Resume.user_id == user.id).order_by(Resume.created_at.desc())
        )
    ]


@router.post("/resumes", status_code=201)
def create_resume(data: ResumeCreate, db: DB, user: CurrentUser) -> dict:
    resume = Resume(user_id=user.id, **data.model_dump())
    db.add(resume)
    commit(db)
    return serialize(resume)


@router.get("/resumes/{resume_id}")
def get_resume(resume_id: str, db: DB, user: CurrentUser) -> dict:
    return serialize(owned(db, Resume, resume_id, user.id))


@router.get("/resumes/{resume_id}/versions")
def list_resume_versions(resume_id: str, db: DB, user: CurrentUser) -> list[dict]:
    owned(db, Resume, resume_id, user.id)
    return [
        serialize(row, exclude=("export_path",))
        for row in db.scalars(
            select(ResumeVersion)
            .where(ResumeVersion.user_id == user.id, ResumeVersion.resume_id == resume_id)
            .order_by(ResumeVersion.version_number.desc())
        )
    ]


@router.post("/resumes/{resume_id}/versions", status_code=201)
def create_resume_version(
    resume_id: str, data: ResumeVersionCreate, db: DB, user: CurrentUser
) -> dict:
    resume = owned(db, Resume, resume_id, user.id)
    if data.parent_version_id:
        parent = owned(db, ResumeVersion, data.parent_version_id, user.id)
        if parent.resume_id != resume.id:
            raise HTTPException(422, "父版本必须属于当前简历")
    content = data.content.model_dump()
    ids = {fact_id for _, _, fact_ids in iter_claims(content) for fact_id in fact_ids}
    facts = list(db.scalars(select(Fact).where(Fact.user_id == user.id, Fact.id.in_(ids))))
    validation = validate_resume_evidence(content, facts)
    if not validation.valid:
        raise HTTPException(422, {"message": "事实引用校验未通过", "errors": validation.errors})
    job, job_version = None, None
    if data.job_id:
        job = owned(db, Job, data.job_id, user.id)
        job_version = db.scalar(
            select(JobVersion)
            .where(JobVersion.user_id == user.id, JobVersion.job_id == job.id)
            .order_by(JobVersion.version_number.desc())
        )
        if job_version is None:
            raise HTTPException(409, "岗位缺少版本快照")
    number = (
        db.scalar(
            select(func.max(ResumeVersion.version_number)).where(
                ResumeVersion.resume_id == resume.id, ResumeVersion.user_id == user.id
            )
        )
        or 0
    ) + 1
    version = ResumeVersion(
        user_id=user.id,
        resume_id=resume.id,
        version_number=number,
        parent_version_id=data.parent_version_id,
        job_id=data.job_id,
        job_version_id=job_version.id if job_version else None,
        language=data.language or resume.language,
        style=data.style or resume.style,
        template=data.template or resume.template,
        content=content,
        facts_snapshot=[fact_snapshot(fact) for fact in facts],
        job_snapshot=job_version.snapshot if job_version else {},
        evidence_status="pending",
    )
    db.add(version)
    try:
        db.flush()
        for evidence in validation.evidence:
            db.add(ResumeEvidence(user_id=user.id, resume_version_id=version.id, **evidence))
        commit(db)
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "版本已经更新，请刷新重试") from None
    return serialize(version, exclude=("export_path",))


def resume_version_for_route(
    db: Session, user_id: str, resume_id: str, version_id: str
) -> ResumeVersion:
    owned(db, Resume, resume_id, user_id)
    version = owned(db, ResumeVersion, version_id, user_id)
    if version.resume_id != resume_id:
        raise HTTPException(404, "简历版本不存在")
    return version


@router.get("/resumes/{resume_id}/versions/{version_id}")
def get_resume_version(resume_id: str, version_id: str, db: DB, user: CurrentUser) -> dict:
    return serialize(
        resume_version_for_route(db, user.id, resume_id, version_id), exclude=("export_path",)
    )


@router.post("/resumes/{resume_id}/versions/{version_id}/approve")
def approve_resume_version(resume_id: str, version_id: str, db: DB, user: CurrentUser) -> dict:
    version = resume_version_for_route(db, user.id, resume_id, version_id)
    if version.evidence_status != "validated":
        raise HTTPException(409, "该版本尚未通过独立语义校验，请先执行 AI 简历验证")
    result = validate_resume_evidence(version.content, version.facts_snapshot)
    if not result.valid:
        raise HTTPException(422, {"message": "事实校验未通过", "errors": result.errors})
    if version.approved_at is None:
        version.approved_at, version.approval_source = utcnow(), "manual"
        commit(db)
    return serialize(version, exclude=("export_path",))


def resume_text(content: dict, language: str = "zh") -> str:
    basics = content.get("basics", {})
    lines = [
        basics.get("name", ""),
        " | ".join(
            str(basics.get(key, "")) for key in ("email", "phone", "location") if basics.get(key)
        ),
    ]
    labels = {
        "summary": "个人简介",
        "education": "教育背景",
        "experience": "项目与经历",
        "skills": "技能",
        "interests": "兴趣与关注",
    }
    if language == "en":
        labels = {
            "summary": "Summary",
            "education": "Education",
            "experience": "Experience",
            "skills": "Skills",
            "interests": "Interests",
        }
    summary = content.get("summary", {}).get("text", "")
    if summary:
        lines.extend(["", labels["summary"], summary])
    for section in ("education", "experience"):
        if content.get(section):
            lines.extend(["", labels[section]])
            for item in content[section]:
                lines.append(item["heading"])
                lines.extend(
                    f"• {bullet['text']}" for bullet in item.get("bullets", []) if bullet["text"]
                )
    for section in ("skills", "interests"):
        if content.get(section):
            lines.extend(["", labels[section]])
            lines.extend(f"• {item['text']}" for item in content[section] if item["text"])
    return "\n".join(lines).strip() + "\n"


def write_pdf(path: Path, version: ResumeVersion) -> None:
    if "STSong-Light" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    compact = version.template == "compact"
    styles = getSampleStyleSheet()
    normal = ParagraphStyle(
        "CareerBody",
        parent=styles["Normal"],
        fontName="STSong-Light",
        fontSize=10 if compact else 10.5,
        leading=15 if compact else 17,
        spaceAfter=4,
        wordWrap="CJK",
    )
    title = ParagraphStyle(
        "CareerName", parent=normal, fontSize=20, leading=25, alignment=TA_CENTER, spaceAfter=8
    )
    heading = ParagraphStyle(
        "CareerSection",
        parent=normal,
        fontSize=12,
        leading=18,
        textColor=colors.HexColor("#1c4b66"),
        spaceBefore=10,
        spaceAfter=5,
    )
    doc = SimpleDocTemplate(
        str(path),
        pagesize=(595.28, 841.89),
        rightMargin=40,
        leftMargin=40,
        topMargin=32 if compact else 42,
        bottomMargin=36,
    )
    content = version.content
    basics = content["basics"]
    story = [Paragraph(escape(basics.get("name", "")), title)]
    contact = " | ".join(
        str(basics.get(key, "")) for key in ("email", "phone", "location") if basics.get(key)
    )
    story.extend(
        [
            Paragraph(escape(contact), normal),
            HRFlowable(width="100%", color=colors.lightgrey),
            Spacer(1, 0.08 * inch),
        ]
    )
    labels = {
        "summary": "个人简介",
        "education": "教育背景",
        "experience": "项目与经历",
        "skills": "技能",
        "interests": "兴趣与关注",
    }
    if version.language == "en":
        labels = {
            "summary": "Summary",
            "education": "Education",
            "experience": "Experience",
            "skills": "Skills",
            "interests": "Interests",
        }
    if content["summary"]["text"]:
        story.extend(
            [
                Paragraph(labels["summary"], heading),
                Paragraph(escape(content["summary"]["text"]).replace("\n", "<br/>"), normal),
            ]
        )
    for section in ("education", "experience"):
        if content[section]:
            story.append(Paragraph(labels[section], heading))
            for item in content[section]:
                story.append(Paragraph(escape(item["heading"]), normal))
                for bullet in item["bullets"]:
                    story.append(
                        Paragraph("• " + escape(bullet["text"]).replace("\n", "<br/>"), normal)
                    )
    for section in ("skills", "interests"):
        if content[section]:
            story.append(Paragraph(labels[section], heading))
            story.extend(
                Paragraph("• " + escape(item["text"]), normal) for item in content[section]
            )
    doc.build(story)


def write_docx(path: Path, version: ResumeVersion) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    document = WordDocument()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10 if version.template == "compact" else 11)
    run_properties = normal.element.get_or_add_rPr()
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "宋体")
    run_properties.append(fonts)
    lines = resume_text(version.content, version.language).splitlines()
    for index, line in enumerate(lines):
        if index == 0:
            document.add_heading(line, level=0)
        elif line.startswith("• "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    document.save(str(path))


def export_resume_version(db: Session, user_id: str, version_id: str, format: str = "pdf") -> Path:
    if format not in {"pdf", "docx", "txt"}:
        raise HTTPException(422, "导出格式支持 pdf/docx/txt")
    version = owned(db, ResumeVersion, version_id, user_id)
    if version.evidence_status != "validated" or version.approved_at is None:
        raise HTTPException(409, "请先完成简历验证与确认后导出")
    fact_ids = {item["id"] for item in version.facts_snapshot}
    facts = list(db.scalars(select(Fact).where(Fact.user_id == user_id, Fact.id.in_(fact_ids))))
    check = validate_resume_evidence(version.content, facts)
    if not check.valid:
        raise HTTPException(422, {"message": "导出前事实校验失败", "errors": check.errors})
    folder = settings.storage_dir.resolve() / user_id / "exports"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{version.id}.{format}"
    if not path.exists():
        with tempfile.NamedTemporaryFile(dir=folder, suffix=f".{format}", delete=False) as handle:
            temporary = Path(handle.name)
        try:
            if format == "pdf":
                write_pdf(temporary, version)
            elif format == "docx":
                write_docx(temporary, version)
            else:
                temporary.write_text(
                    resume_text(version.content, version.language), encoding="utf-8"
                )
            temporary.replace(path)
        finally:
            temporary.unlink(missing_ok=True)
    version.export_status, version.export_path = "ready", str(path)
    db.flush()
    return path


@router.post("/resumes/{resume_id}/versions/{version_id}/export")
def export_version(
    resume_id: str,
    version_id: str,
    db: DB,
    user: CurrentUser,
    format: str = Query(default="pdf", pattern="^(pdf|docx|txt)$"),
) -> dict:
    resume_version_for_route(db, user.id, resume_id, version_id)
    export_resume_version(db, user.id, version_id, format)
    commit(db)
    return {
        "status": "ready",
        "format": format,
        "download_url": f"/api/v1/resumes/{resume_id}/versions/{version_id}/download?format={format}",
    }


@router.get("/resumes/{resume_id}/versions/{version_id}/download")
def download_version(
    resume_id: str,
    version_id: str,
    db: DB,
    user: CurrentUser,
    format: str = Query(default="pdf", pattern="^(pdf|docx|txt)$"),
) -> FileResponse:
    resume_version_for_route(db, user.id, resume_id, version_id)
    path = export_resume_version(db, user.id, version_id, format)
    commit(db)
    media_types = {
        "pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "txt": "text/plain; charset=utf-8",
    }
    return FileResponse(
        path, media_type=media_types[format], filename=f"resume-{version_id[:8]}.{format}"
    )


@router.get("/applications")
def list_applications(db: DB, user: CurrentUser) -> list[dict]:
    applications = db.scalars(
        select(Application)
        .where(Application.user_id == user.id)
        .order_by(Application.created_at.desc())
    ).all()
    job_ids = {application.job_id for application in applications}
    jobs = {
        job.id: job
        for job in db.scalars(select(Job).where(Job.user_id == user.id, Job.id.in_(job_ids)))
    }
    return [
        {
            **serialize(application),
            "company": jobs[application.job_id].company,
            "job_title": jobs[application.job_id].title,
        }
        for application in applications
    ]


@router.post("/applications", status_code=201)
def create_application(data: ApplicationCreate, db: DB, user: CurrentUser) -> dict:
    owned(db, Job, data.job_id, user.id)
    if data.resume_version_id:
        owned(db, ResumeVersion, data.resume_version_id, user.id)
    application = Application(user_id=user.id, **data.model_dump())
    db.add(application)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "该岗位投递轮次已存在，请使用新的 attempt_no") from None
    db.add(
        ApplicationEvent(
            user_id=user.id,
            application_id=application.id,
            sequence=1,
            from_status=None,
            to_status=application.status,
            notes="用户手动建立投递记录",
        )
    )
    commit(db)
    return serialize(application)


@router.patch("/applications/{application_id}")
def patch_application(
    application_id: str, data: ApplicationUpdate, db: DB, user: CurrentUser
) -> dict:
    application = owned(db, Application, application_id, user.id)
    changes = data.model_dump(exclude_unset=True, exclude={"version"})
    nullable = {"applied_at", "resume_version_id"}
    if any(value is None and key not in nullable for key, value in changes.items()):
        raise HTTPException(422, "状态和文本字段不能设为 null")
    if data.resume_version_id:
        owned(db, ResumeVersion, data.resume_version_id, user.id)
    next_status = changes.get("status", application.status)
    if (
        application.status in {"offer", "rejected", "withdrawn"}
        and next_status != application.status
    ):
        raise HTTPException(409, "该投递轮次已结束；再次投递请建立新的 attempt_no")
    previous_status = application.status
    result = db.execute(
        update(Application)
        .where(
            Application.id == application.id,
            Application.user_id == user.id,
            Application.version == data.version,
        )
        .values(**changes, version=data.version + 1, updated_at=utcnow())
    )
    if result.rowcount != 1:
        db.rollback()
        raise HTTPException(409, "投递记录已更新，请刷新后重试")
    db.add(
        ApplicationEvent(
            user_id=user.id,
            application_id=application.id,
            sequence=data.version + 1,
            from_status=previous_status,
            to_status=next_status,
            notes=changes.get("notes", "用户手动更新"),
        )
    )
    commit(db)
    db.refresh(application)
    return serialize(application)


@router.get("/applications/{application_id}/events")
def list_application_events(application_id: str, db: DB, user: CurrentUser) -> list[dict]:
    owned(db, Application, application_id, user.id)
    return [
        serialize(event)
        for event in db.scalars(
            select(ApplicationEvent)
            .where(
                ApplicationEvent.user_id == user.id,
                ApplicationEvent.application_id == application_id,
            )
            .order_by(ApplicationEvent.sequence)
        )
    ]


SOURCE_ENTRIES = [
    ("BOSS 直聘", "website", "https://www.zhipin.com/", "仅登记入口，未获授权自动投递"),
    ("LinkedIn", "website", "https://www.linkedin.com/jobs/", "仅登记入口"),
    ("实习僧", "website", "https://www.shixiseng.com/", "仅登记入口"),
    ("51job", "website", "https://www.51job.com/", "仅登记入口"),
    ("猎聘", "website", "https://www.liepin.com/", "仅登记入口"),
    ("Offer Cat", "website", "https://offer-cat.com/", "用户提供入口，数据未同步"),
    ("求职方舟", "website", "https://www.qiuzhifangzhou.com/", "用户提供入口，数据未同步"),
    ("OfferShow", "website", "https://www.offershow.cn/", "用户提供入口，数据未同步"),
    ("OfferShow 小程序", "mini_program", "", "小程序入口待用户补充，未集成"),
    ("牛客汇总", "wechat", "", "用户提供公众号名称；尚未核验官方账号"),
    ("牛客校招日程", "website", "https://www.nowcoder.com/jobs/school/schedule", "仅登记入口"),
    (
        "【秋招/春招/实习】汇总表-婉清学姐",
        "feishu",
        "https://yal2at57cvq.feishu.cn/base/GtSLbyyR3aCENOsJYC6cdlsVnih?table=tblH4au5rnBcqHgJ&view=vew8PFC7nG",
        "用户提供飞书表格；同步需配置读取授权及资源允许列表",
    ),
    (
        "互联派名企校招2",
        "spreadsheet",
        "https://docs.qq.com/sheet/DTHFWSWpyeXVxSUJw?tab=BB08J2",
        "官方校招/名企直推/名企招聘/外企/简历模板；用户可导出 CSV/XLSX 导入，未集成在线抓取",
    ),
]


@router.post("/demo/seed")
def seed_demo(db: DB, user: CurrentUser) -> dict:
    existing = db.scalar(
        select(Fact.id).where(Fact.user_id == user.id, Fact.title == "[虚构演示] 教学资料整理项目")
    )
    if existing:
        return {"seeded": False, "message": "当前用户已建立虚构演示资料；不会重复添加"}
    sources = []
    for name, kind, url, notes in SOURCE_ENTRIES:
        source = db.scalar(select(Source).where(Source.user_id == user.id, Source.name == name))
        if source is None:
            source = Source(user_id=user.id, name=name, kind=kind, url=url, notes=notes)
            db.add(source)
            db.flush()
        sources.append(source)
    fact = Fact(
        user_id=user.id,
        kind="project",
        title="[虚构演示] 教学资料整理项目",
        content="虚构体验数据：协助教学项目整理课程资料，使用表格建立资料目录并检查重复项。",
        tags=["虚构演示", "资料管理", "表格"],
        source_text="由演示种子建立，不代表用户真实经历。",
        confirmed=False,
    )
    interest = Fact(
        user_id=user.id,
        kind="interest",
        title="[虚构演示] 数字产品观察",
        content="虚构体验数据：长期关注搜索和协作工具，记录公开产品的用户体验观察。",
        tags=["虚构演示"],
        source_text="由演示种子建立，不代表用户真实经历。",
    )
    db.add_all([fact, interest])
    create_job_record(
        db,
        user.id,
        JobCreate(
            company="[虚构演示] 星桥教育",
            title="[虚构演示] 教学运营助理",
            location="上海",
            industry="教育",
            jd_text="虚构岗位，仅用于体验：整理课程资料，维护教学项目表格，协助课程运营；重视资料管理、沟通和表格工具使用。",
            source_id=sources[0].id,
            source_url="https://example.com/fictional-career-job",
        ),
        "demo",
    )
    db.add(Resume(user_id=user.id, name="[虚构演示] 简历工作稿", language="zh"))
    commit(db)
    return {
        "seeded": True,
        "fictional": True,
        "sources_count": len(sources),
        "message": "已添加明确标注的虚构岗位和事实；事实尚未确认。请用自己的真实经历替换后生成简历。",
    }
